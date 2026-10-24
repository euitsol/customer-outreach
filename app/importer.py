"""Bulk import / export of contacts and their per-person email content (CSV or XLSX)."""
import csv
import io
import json
import os
import re
import secrets
import tempfile
import time
from datetime import date, datetime

from dateutil import parser as dateparser

from app import settings as app_settings
from app.crm import activities as activities_svc, audit, people as people_svc
from app.merge import slug
from app.models import Activity, CampaignStep, Enrollment
from app.services import EMAIL_RE, enroll, get_or_create_message, recompute_enrollment, steps_of

IMPORT_DIR = os.path.join(tempfile.gettempdir(), "outreach_imports")
MAX_STEP_INDEX = 12

BASE_TARGETS = [
    ("ignore", "Ignore this column"),
    ("email", "Email (required)"),
    ("first_name", "First name"),
    ("last_name", "Last name"),
    ("full_name", "Full name (split)"),
    ("company", "Company"),
    ("job_title", "Job title"),
    ("company_website", "Company website"),
    ("phone", "Phone"),
    ("priority", "Priority"),
    ("tags", "Tags"),
    ("notes", "Notes"),
    ("send", "Send? (Send / Do not send / Verify)"),
    ("custom", "Custom merge field"),
]


def step_targets(n_steps):
    out = []
    for n in range(1, max(n_steps, 4) + 1):
        out += [(f"step{n}_subject", f"Stage {n}: subject"),
                (f"step{n}_body", f"Stage {n}: email body"),
                (f"step{n}_date", f"Stage {n}: send date")]
    return out


# ------------------------------------------------------------------ reading files

def _sniff_csv(text):
    sample = text[:4000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        dialect.doublequote = True
        return dialect
    except csv.Error:
        return csv.excel


def read_table(filename, raw):
    """Return (headers, rows) where rows are lists of strings. Detects the header row."""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        grid = [["" if v is None else (v.date().isoformat() if isinstance(v, datetime) else str(v))
                 for v in row] for row in ws.iter_rows(values_only=True)]
    else:
        text = raw.decode("utf-8-sig", errors="replace")
        grid = list(csv.reader(io.StringIO(text), _sniff_csv(text)))
    grid = [r for r in grid if any((c or "").strip() for c in r)]
    if not grid:
        raise ValueError("The file is empty.")
    header_idx = None
    for i, row in enumerate(grid[:25]):
        if any(re.fullmatch(r"\s*e-?mail( address)?\s*", c or "", re.I) or re.fullmatch(r"\s*email\s*\*?\s*", c or "", re.I)
               for c in row):
            header_idx = i
            break
    if header_idx is None:
        # fall back: first row whose any cell is exactly something containing 'email'
        for i, row in enumerate(grid[:25]):
            if any("email" in (c or "").lower() and "@" not in (c or "") for c in row):
                header_idx = i
                break
    if header_idx is None:
        raise ValueError("Could not find a header row with an 'Email' column.")
    headers, seen = [], {}
    for h in grid[header_idx]:
        h = (h or "").strip() or "column"
        seen[h] = seen.get(h, 0) + 1
        headers.append(h if seen[h] == 1 else f"{h} ({seen[h]})")
    width = len(headers)
    rows = [(r + [""] * width)[:width] for r in grid[header_idx + 1:]]
    return headers, rows


def save_upload(filename, raw):
    os.makedirs(IMPORT_DIR, exist_ok=True)
    cutoff = time.time() - 86400
    for f in os.listdir(IMPORT_DIR):
        p = os.path.join(IMPORT_DIR, f)
        try:
            if os.path.getmtime(p) < cutoff:
                os.remove(p)
        except OSError:
            pass
    token = secrets.token_hex(12)
    with open(os.path.join(IMPORT_DIR, token + ".bin"), "wb") as fh:
        fh.write(raw)
    with open(os.path.join(IMPORT_DIR, token + ".json"), "w") as fh:
        json.dump({"filename": filename}, fh)
    return token


def load_upload(token):
    if not re.fullmatch(r"[0-9a-f]{24}", token or ""):
        raise ValueError("Upload expired. Please upload the file again.")
    try:
        with open(os.path.join(IMPORT_DIR, token + ".bin"), "rb") as fh:
            raw = fh.read()
        with open(os.path.join(IMPORT_DIR, token + ".json")) as fh:
            name = json.load(fh)["filename"]
    except OSError:
        raise ValueError("Upload expired. Please upload the file again.")
    return name, raw


# ------------------------------------------------------------------ column mapping

DAY_DATE = re.compile(r"^(mon|tue|wed|thu|fri|sat|sun)[a-z]*\s+\d{1,2}\s+[a-z]{3}", re.I)


def guess_target(header):
    h = (header or "").strip()
    c = re.sub(r"[^a-z0-9]", "", re.sub(r"\(\d+\)$", "", h.lower()))
    if DAY_DATE.match(h):
        return "step1_date"
    if c in ("email", "emailaddress", "mail", "emailaddr"):
        return "email"
    m = re.fullmatch(r"(?:email|mail|message|body|stage|step)(\d+)(body|text|message)?", c)
    if m:
        return f"step{int(m.group(1))}_body"
    m = re.fullmatch(r"(?:email|mail|stage|step)(\d+)(subject|subjectline)", c)
    if m:
        return f"step{int(m.group(1))}_subject"
    m = re.fullmatch(r"(?:email|mail|stage|step)(\d+)(date|senddate|sendon)", c)
    if m:
        return f"step{int(m.group(1))}_date"
    m = re.fullmatch(r"followup(\d+)(body|text|message|email)?", c)
    if m:
        return f"step{int(m.group(1)) + 1}_body"
    m = re.fullmatch(r"followup(\d+)(subject|subjectline)", c)
    if m:
        return f"step{int(m.group(1)) + 1}_subject"
    m = re.fullmatch(r"followup(\d+)(date|senddate|sendon)", c)
    if m:
        return f"step{int(m.group(1)) + 1}_date"
    simple = {
        "subject": "step1_subject", "subjectline": "step1_subject",
        "firstname": "first_name", "first": "first_name", "fname": "first_name", "givenname": "first_name",
        "lastname": "last_name", "last": "last_name", "surname": "last_name", "lname": "last_name",
        "name": "full_name", "fullname": "full_name", "contactname": "full_name",
        "company": "company", "companyname": "company", "business": "company",
        "organisation": "company", "organization": "company", "account": "company",
        "title": "job_title", "jobtitle": "job_title", "position": "job_title", "role": "job_title",
        "website": "company_website", "companywebsite": "company_website", "web": "company_website",
        "url": "company_website", "domain": "company_website",
        "phone": "phone", "telephone": "phone", "mobile": "phone", "phonenumber": "phone",
        "priority": "priority", "tags": "tags", "tag": "tags",
        "notes": "notes", "note": "notes", "comments": "notes",
        "send": "send", "sendstatus": "send",
        "status": "ignore", "no": "ignore", "id": "ignore",
    }
    if c in simple:
        return simple[c]
    if c.startswith("notes") or c.startswith("comments"):
        return "notes"
    return "custom" if c else "ignore"


def guess_mapping(headers):
    mapping = [guess_target(h) for h in headers]
    seen = set()
    for i, t in enumerate(mapping):
        if t not in ("ignore", "custom"):
            if t in seen:
                mapping[i] = "custom" if not t.startswith("step") else "ignore"
            seen.add(t)
    if "email" not in mapping:
        for i, h in enumerate(headers):
            if "email" in h.lower() and mapping[i] in ("custom", "ignore"):
                mapping[i] = "email"
                break
    return mapping


# ------------------------------------------------------------------ parsing rows

def parse_date_cell(value, today=None):
    v = (value or "").strip()
    if not v:
        return None
    today = today or date.today()
    try:
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", v):
            return datetime.strptime(v, "%Y-%m-%d").date()
        d = dateparser.parse(v, fuzzy=True, dayfirst=True, default=datetime(today.year, 1, 1)).date()
    except (ValueError, OverflowError):
        return None
    has_year = bool(re.search(r"\d{4}", v))
    if not has_year and (today - d).days > 183:
        d = d.replace(year=d.year + 1)
    return d


def normalise_send(value):
    v = (value or "").strip().lower()
    if v in ("", "send", "yes", "y", "true", "1"):
        return "send"
    if v.startswith("verify"):
        return "verify"
    if v.startswith(("do not", "dont", "don't", "no", "skip", "false", "0", "pause")):
        return "skip"
    return None


def parse_rows(rows, headers, mapping, custom_as_fields=True):
    """Turn raw rows into normalised dicts. Returns (parsed, errors)."""
    parsed, errors = [], []
    custom_names = {i: slug(h) for i, h in enumerate(headers) if mapping[i] == "custom"}
    for n, row in enumerate(rows, start=2):
        rec = {"row": n, "custom": {}, "steps": {}, "tags": ""}
        for i, target in enumerate(mapping):
            val = (row[i] or "").strip() if i < len(row) else ""
            if target == "ignore" or not val:
                continue
            if target == "custom":
                if custom_as_fields and custom_names.get(i):
                    rec["custom"][custom_names[i]] = val
            elif target.startswith("step"):
                idx, _, kind = target[4:].partition("_")
                st = rec["steps"].setdefault(int(idx), {})
                st[kind] = val
            elif target == "full_name":
                first, _, last = val.partition(" ")
                rec.setdefault("first_name", first)
                rec.setdefault("last_name", last)
            else:
                rec[target] = val
        email = (rec.get("email") or "").strip().lower()
        if not email:
            if any((c or "").strip() for c in row):
                errors.append({"row": n, "email": "", "error": "missing email"})
            continue
        if not EMAIL_RE.match(email):
            errors.append({"row": n, "email": email, "error": "invalid email format"})
            continue
        rec["email"] = email
        send = normalise_send(rec.get("send"))
        if send is None:
            errors.append({"row": n, "email": email, "error": f"unrecognised Send? value: {rec.get('send')}"})
            continue
        rec["send"] = send
        for idx, st in rec["steps"].items():
            if "date" in st:
                d = parse_date_cell(st["date"])
                if d is None:
                    errors.append({"row": n, "email": email, "error": f"bad date for stage {idx}: {st['date']}"})
                    st.pop("date")
                else:
                    st["date"] = d
        parsed.append(rec)
    return parsed, errors


def analyse(session, parsed, campaign, mode="upsert"):
    """Dry run: classify every row without writing."""
    existing = {}
    if campaign is not None and campaign.id:
        existing = {e.email.lower(): e for e in session.query(Enrollment).filter(Enrollment.campaign_id == campaign.id)}
    seen, plan = set(), []
    stats = {"creates": 0, "updates": 0, "skip_existing": 0, "skip_missing": 0, "skip_dnr": 0,
             "dupes_in_file": 0, "customised": 0}
    for rec in parsed:
        e = rec["email"]
        if e in seen:
            stats["dupes_in_file"] += 1
            continue
        seen.add(e)
        exists = e in existing
        if rec["send"] == "skip" and not exists:
            stats["skip_dnr"] += 1
            continue
        if mode == "create_only" and exists:
            stats["skip_existing"] += 1
            continue
        if mode == "update_only" and not exists:
            stats["skip_missing"] += 1
            continue
        stats["updates" if exists else "creates"] += 1
        if any(s.get("body") or s.get("subject") for s in rec["steps"].values()):
            stats["customised"] += 1
        plan.append((rec, existing.get(e)))
    return plan, stats


def apply_plan(session, campaign, plan, settings=None):
    steps = {s.position: s for s in steps_of(session, campaign.id)}
    max_idx = max([i for rec, _ in plan for i in rec["steps"]] or [0])
    if max_idx > MAX_STEP_INDEX:
        raise ValueError(f"Too many stages in file (max {MAX_STEP_INDEX}).")
    for idx in range(1, max_idx + 1):
        if idx not in steps:
            st = CampaignStep(campaign_id=campaign.id, position=idx,
                              name="Email 1" if idx == 1 else f"Follow-up {idx - 1}",
                              delay_hours=0 if idx == 1 else 72)
            session.add(st)
            session.flush()
            steps[idx] = st
    ordered = [steps[k] for k in sorted(steps)]

    created = updated = locked = 0
    actor = audit.Actor("system", label="import")
    for rec, enrollment in plan:
        is_new = enrollment is None
        # The file is the newest information we have: it refreshes whatever it provides.
        person, _ = people_svc.upsert(
            session, email=rec["email"], phone=rec.get("phone"), first_name=rec.get("first_name"),
            last_name=rec.get("last_name"), job_title=rec.get("job_title"), company_name=rec.get("company"),
            company_website=rec.get("company_website"), source="import", overwrite=True, actor=actor)
        if is_new:
            enrollment, _ = enroll(session, campaign, person, email=rec["email"], steps=ordered)
        if rec.get("priority"):
            enrollment.priority = rec["priority"]
        if rec["custom"]:
            merged = enrollment.custom_fields
            merged.update(rec["custom"])
            enrollment.custom_fields = merged
        if rec["send"] == "verify":
            enrollment.needs_verification = True
        elif rec["send"] == "skip":
            enrollment.send_flag = False
        if rec.get("tags"):
            people_svc.add_tags(session, person, rec["tags"])
        note = (rec.get("notes") or "").strip()
        if note and not session.query(Activity).filter(Activity.person_id == person.id, Activity.kind == "note",
                                                       Activity.body == note).first():
            activities_svc.log(session, kind="note", subject="Note (import)", body=note, person_id=person.id)
        for idx, data in rec["steps"].items():
            if not (data.get("subject") or data.get("body") or data.get("date")):
                continue
            msg = get_or_create_message(session, enrollment, steps[idx])
            if msg.status == "sent":
                locked += 1
                continue
            if data.get("subject"):
                msg.subject = data["subject"]
            if data.get("body"):
                msg.body = data["body"]
            if data.get("date"):
                midnight = datetime.combine(data["date"], datetime.min.time())
                msg.due_at = app_settings.local_to_utc(midnight, settings) if settings else midnight
        session.flush()
        recompute_enrollment(session, enrollment, ordered, campaign)
        created += is_new
        updated += not is_new
    return created, updated, locked


# ------------------------------------------------------------------ export / templates

CORE_COLUMNS = ["email", "first_name", "last_name", "job_title", "company", "company_website", "phone", "priority",
                "send", "tags"]


def enrollment_status_label(c):
    if c.bounced:
        return "bounced"
    if c.unsubscribed:
        return "unsubscribed"
    if c.replied:
        return "replied"
    if c.needs_verification:
        return "verify"
    if not c.send_flag:
        return "paused"
    return "ready" if not c.stage else f"stage {c.stage} sent"


def export_csv(session, enrollments, steps):
    """Round-trippable export: same column names the importer understands."""
    contacts = enrollments
    custom_keys = []
    for c in contacts:
        for k in c.custom_fields:
            if k not in custom_keys:
                custom_keys.append(k)
    n = max(len(steps), 1)
    step_cols = []
    for i in range(1, n + 1):
        step_cols += [f"stage{i}_subject", f"stage{i}_body", f"stage{i}_date"]
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(CORE_COLUMNS + step_cols + custom_keys + ["current_stage", "status"])
    step_by_pos = {s.position: s for s in steps}
    for c in contacts:
        msgs = {m.position: m for m in c.messages}
        send_val = "Verify email first" if c.needs_verification else ("Send" if c.send_flag else "Do not send")
        p = c.person
        company = p.company if p else None
        row = [c.email, c.first_name, c.last_name, p.job_title if p else "", c.company,
               company.website if company else "", c.phone, c.priority, send_val,
               ",".join(sorted(t.name for t in c.tags))]
        for i in range(1, n + 1):
            m = msgs.get(i)
            row += [m.subject if m else "", m.body if m else "", m.due_at.date().isoformat() if m and m.due_at else ""]
        row += [c.custom_fields.get(k, "") for k in custom_keys]
        row += [c.stage or 0, enrollment_status_label(c)]
        w.writerow(row)
    return buf.getvalue()


def template_csv(steps):
    n = max(len(steps), 3)
    cols = ["email", "first_name", "last_name", "job_title", "company", "company_website", "phone", "priority", "send",
            "tags", "notes"]
    for i in range(1, n + 1):
        cols += [f"stage{i}_subject", f"stage{i}_body", f"stage{i}_date"]
    cols += ["what_we_found"]
    sample = ["jane@acmehotel.ie", "Jane", "Doe", "General Manager", "Acme Hotel", "https://acmehotel.ie",
              "+353 1 000 0000", "High", "Send", "hotels,dublin", "Decision maker"]
    for i in range(1, n + 1):
        if i == 1:
            sample += ["Your booking link isn't working", "Hi Jane,\n\nI noticed your booking button...\n\nBest regards,", ""]
        else:
            sample += ["", f"Hi Jane,\n\nFollowing up on my note (stage {i})...\n\nBest regards,", ""]
    sample += ["Booking button on the mobile site is broken"]
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(cols)
    w.writerow(sample)
    return buf.getvalue()
